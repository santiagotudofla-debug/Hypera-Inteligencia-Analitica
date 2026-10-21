import streamlit as st
import pandas as pd
import io
from datetime import datetime
from fpdf import FPDF, XPos, YPos
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from utils import (
    carregar_demonstrativos_cvm_real,
    carregar_info_ticker,
    obter_dre_cvm_estruturado,
    calcular_altman_z_score,
    sanitizar_texto_pdf
)

st.title("📥 Central de Exportação e Relatórios Profissionais")
st.markdown("Baixe relatórios executivos ricos em dados, contendo matrizes do DRE, Valuation e ESG formatados para sua apresentação.")

# Carregar dados principais via cache centralizado
ano_atual = datetime.now().year
df_cvm_bruto = carregar_demonstrativos_cvm_real(ano_atual)
if df_cvm_bruto.empty:
    df_cvm_bruto = carregar_demonstrativos_cvm_real(ano_atual - 1)

info = carregar_info_ticker("HYPE3.SA")
preco = float(info.get('currentPrice', info.get('regularMarketPrice', 0.0)) or 0.0)
roe = float(info.get('returnOnEquity', 0.0) or 0.0) * 100
dy = float(info.get('dividendYield', 0.0) or 0.0) * 100
margem_liq = float(info.get('profitMargins', 0.0) or 0.0) * 100
divida = float(info.get('totalDebt', 0.0) or 0.0)
receita = float(info.get('totalRevenue', 0.0) or 0.0)

# Obter DRE estruturado, filtrado e sem duplicações temporais
df_dre_estruturado = obter_dre_cvm_estruturado(df_cvm_bruto)
df_dre = df_dre_estruturado.head(10)

# Cálculo real de Z-Score
z_analise = calcular_altman_z_score(info)
z_val = z_analise["z_score"]
z_zona = z_analise["zona"]

# --- 1. Exportação CSV ---
st.markdown("### 📊 1. Base de Dados Contábil (CSV)")
st.write("Baixe a planilha estruturada com os demonstrativos oficiais da CVM desduplicados para análises customizadas no Excel.")

@st.cache_data
def convert_df(df):
    return df.to_csv(index=False).encode('utf-8')

csv_data = convert_df(df_dre_estruturado)
st.download_button(
    label="Baixar Planilha CVM (.csv)",
    data=csv_data,
    file_name=f'hypera_dre_oficial_{ano_atual}.csv',
    mime='text/csv',
)

st.markdown("---")

# --- 2. Geração de PowerPoint (PPTX) ---
st.markdown("### 📽️ 2. Apresentação Executiva Completa (PowerPoint)")
st.write("Gera múltiplos slides profissionais cobrindo ESG, DRE, Risco e Valuation com dados matemáticos reais.")

def generate_pptx():
    prs = Presentation()
    
    # 1. Slide de Capa
    slide_capa = prs.slides.add_slide(prs.slide_layouts[0])
    slide_capa.shapes.title.text = "Relatório Executivo Analítico\nHypera Pharma (HYPE3)"
    slide_capa.placeholders[1].text = f"Gerado pelo Motor de Inteligência Financeira\nData: {datetime.now().strftime('%d/%m/%Y')}"

    # 2. Slide Visão Corporativa e ESG
    slide_esg = prs.slides.add_slide(prs.slide_layouts[1])
    slide_esg.shapes.title.text = "Visão Corporativa e Sustentabilidade (ESG)"
    tf_esg = slide_esg.shapes.placeholders[1].text_frame
    tf_esg.text = "Destaques do Negócio:"
    
    p = tf_esg.add_paragraph()
    p.text = "Setor B3: Saúde (Produtos Farmacêuticos)"
    p.level = 1
    p = tf_esg.add_paragraph()
    p.text = "Governança: Novo Mercado (100% Tag Along)"
    p.level = 1
    p = tf_esg.add_paragraph()
    p.text = "Metas ESG Ativas: ODS 3 (Saúde) e Eficiência Hídrica"
    p.level = 1

    # 3. Slide de Tabela DRE
    slide_dre = prs.slides.add_slide(prs.slide_layouts[5]) # Titulo apenas
    slide_dre.shapes.title.text = "Demonstrativo de Resultados (DRE Oficial CVM)"
    
    rows = min(6, len(df_dre) + 1)
    cols = 2
    left = Inches(1)
    top = Inches(2)
    width = Inches(8)
    height = Inches(0.8)
    
    table = slide_dre.shapes.add_table(rows, cols, left, top, width, height).table
    table.columns[0].width = Inches(5.5)
    table.columns[1].width = Inches(2.5)
    
    table.cell(0, 0).text = "Conta Contábil"
    table.cell(0, 1).text = "Valor Declarado (R$)"
    
    for i in range(1, rows):
        row_data = df_dre.iloc[i-1]
        desc = f"{row_data.get('CD_CONTA', '')} - {row_data.get('DS_CONTA', '')}"
        table.cell(i, 0).text = desc[:60]
        valor = row_data.get("VL_CONTA", 0)
        table.cell(i, 1).text = f"{valor:,.0f}"

    # 4. Slide de Valuation e Risco
    slide_val = prs.slides.add_slide(prs.slide_layouts[1])
    slide_val.shapes.title.text = "Valuation e Indicadores de Risco"
    tf_val = slide_val.shapes.placeholders[1].text_frame
    tf_val.text = "Métricas em Tempo Real (B3 & Modelo Altman Z''):"
    
    p = tf_val.add_paragraph()
    p.text = f"Cotação Atual: R$ {preco:,.2f}"
    p.level = 1
    p = tf_val.add_paragraph()
    p.text = f"ROE (Rentabilidade Patrimonial): {roe:.1f}%"
    p.level = 1
    p = tf_val.add_paragraph()
    p.text = f"Margem Líquida (Eficiência): {margem_liq:.1f}%"
    p.level = 1
    p = tf_val.add_paragraph()
    p.text = f"Dividend Yield: {dy:.1f}%"
    p.level = 1
    p = tf_val.add_paragraph()
    p.text = f"Altman Z''-Score Calculado: {z_val:.2f} ({z_zona})"
    p.level = 1
    
    pptx_io = io.BytesIO()
    prs.save(pptx_io)
    pptx_io.seek(0)
    return pptx_io

pptx_file = generate_pptx()
st.download_button(
    label="Baixar Apresentação Profissional (.pptx)",
    data=pptx_file,
    file_name=f'Relatorio_Master_HYPE3_{ano_atual}.pptx',
    mime='application/vnd.openxmlformats-officedocument.presentationml.presentation'
)

st.markdown("---")

# --- 3. Geração de PDF ---
st.markdown("### 📄 3. Relatório Analítico em Documento (PDF)")
st.write("Um documento corporativo detalhado com matrizes de dados, pareces técnicos e governança.")

def generate_pdf():
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    
    # Capa / Cabeçalho
    pdf.add_page()
    pdf.set_font("Helvetica", 'B', 18)
    pdf.set_text_color(0, 51, 102) # Azul Corporativo
    pdf.cell(0, 15, sanitizar_texto_pdf("Relatório de Inteligência Financeira: Hypera S.A."), new_x=XPos.LMARGIN, new_y=YPos.NEXT, align='C')
    pdf.set_font("Helvetica", 'I', 11)
    pdf.set_text_color(100, 100, 100)
    pdf.cell(0, 8, sanitizar_texto_pdf(f"Ticker: HYPE3 | Extração: {datetime.now().strftime('%d/%m/%Y')} | Fonte: CVM / B3"), new_x=XPos.LMARGIN, new_y=YPos.NEXT, align='C')
    pdf.ln(8)
    
    # 1. Resumo Corporativo e ESG
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Helvetica", 'B', 13)
    pdf.cell(0, 8, sanitizar_texto_pdf("1. Visão Corporativa e ESG"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", '', 10)
    texto_esg = sanitizar_texto_pdf(
        "A Hypera Pharma atua no segmento de produtos farmacêuticos e saúde. "
        "A empresa está listada no Novo Mercado da B3, garantindo o nível máximo de "
        "governança corporativa (100% Tag Along) e passa por auditoria independente "
        "com registro regular na CVM. No contexto ESG, destaca-se pelo engajamento com "
        "a agenda ODS (Saúde e Bem-estar)."
    )
    pdf.multi_cell(0, 5, texto_esg)
    pdf.ln(6)
    
    # 2. Indicadores de Mercado (Valuation)
    pdf.set_font("Helvetica", 'B', 13)
    pdf.cell(0, 8, sanitizar_texto_pdf("2. Indicadores de Mercado e Valuation em Tempo Real"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", '', 10)
    
    # Tabela de métricas
    pdf.set_fill_color(240, 240, 240)
    pdf.set_font("Helvetica", 'B', 10)
    pdf.cell(95, 7, sanitizar_texto_pdf("Métrica Fundamentalista"), border=1, fill=True)
    pdf.cell(95, 7, sanitizar_texto_pdf("Valor Atual"), border=1, fill=True, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    
    pdf.set_font("Helvetica", '', 10)
    metricas = [
        ("Cotação Atual", f"R$ {preco:,.2f}"),
        ("Receita Total Estimada", f"R$ {receita:,.2f}"),
        ("Dívida Total", f"R$ {divida:,.2f}"),
        ("Margem Líquida", f"{margem_liq:.1f}%"),
        ("Retorno sobre Patrimônio (ROE)", f"{roe:.1f}%"),
        ("Dividend Yield", f"{dy:.1f}%"),
        ("Altman Z''-Score Calculado", f"{z_val:.2f} ({z_zona})")
    ]
    for label, val in metricas:
        pdf.cell(95, 7, sanitizar_texto_pdf(label), border=1)
        pdf.cell(95, 7, sanitizar_texto_pdf(val), border=1, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    
    pdf.ln(8)
    
    # 3. Tabela do DRE (CVM)
    pdf.set_font("Helvetica", 'B', 13)
    pdf.cell(0, 8, sanitizar_texto_pdf("3. Resumo da DRE (Base Oficial CVM - Último Período)"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", 'I', 9)
    pdf.cell(0, 6, sanitizar_texto_pdf("Contas contábeis desduplicadas extraídas dos demonstrativos ITR oficiais:"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    
    pdf.set_font("Helvetica", 'B', 9)
    pdf.set_fill_color(0, 51, 102)
    pdf.set_text_color(255, 255, 255)
    pdf.cell(140, 7, sanitizar_texto_pdf("Descrição da Conta (CVM)"), border=1, fill=True)
    pdf.cell(50, 7, sanitizar_texto_pdf("Valor Declarado (R$)"), border=1, fill=True, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align='R')
    
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Helvetica", '', 9)
    
    for idx, row in df_dre.head(8).iterrows():
        cd = str(row.get("CD_CONTA", ""))
        ds = str(row.get("DS_CONTA", "N/D"))
        descricao = sanitizar_texto_pdf(f"{cd} - {ds}")[:65]
        valor = float(row.get("VL_CONTA", 0) or 0)
        pdf.cell(140, 6, descricao, border=1)
        pdf.cell(50, 6, f"{valor:,.0f}", border=1, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align='R')
        
    pdf.ln(8)
    
    # 4. Conclusão e Risco
    pdf.set_font("Helvetica", 'B', 12)
    pdf.cell(0, 8, sanitizar_texto_pdf("4. Parecer de Risco Analítico e Solvência"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", '', 10)
    texto_parecer = sanitizar_texto_pdf(
        f"Com base na modelagem estatística de risco Altman Z''-Score ({z_val:.2f}), a Hypera Pharma "
        f"encontra-se classificada na {z_zona}. A estrutura de capital apresenta alavancagem em nível compatível "
        "com a geração operacional de caixa, demonstrando previsibilidade para atendimento de passivos e sustentabilidade na remuneração acionária."
    )
    pdf.multi_cell(0, 5, texto_parecer)
    
    # Rodapé final
    pdf.set_y(-22)
    pdf.set_font('Helvetica', 'I', 8)
    pdf.cell(0, 8, sanitizar_texto_pdf('Documento gerado automaticamente pela Plataforma HYPE3 Inteligência Financeira.'), 0, 0, 'C')

    pdf_bytes = pdf.output()
    return bytes(pdf_bytes)

pdf_file = generate_pdf()
st.download_button(
    label="Baixar Relatório Executivo Completo (.pdf)",
    data=pdf_file,
    file_name=f'Relatorio_Executivo_HYPE3_{ano_atual}.pdf',
    mime='application/pdf'
)

